from flask import Flask, request, jsonify, render_template
from flask_cors import CORS
from docx import Document
from pptx import Presentation
import fitz
import requests

app = Flask(__name__)
CORS(app)
def extract_pdf(file):

    text = ""

    pdf = fitz.open(stream=file.read(), filetype="pdf")

    for page in pdf:
        text += page.get_text()

    return text


def summarize_ai(text):

    response = requests.post(

        url="https://openrouter.ai/api/v1/chat/completions",

        headers={

            "Authorization": f"Bearer {OPENROUTER_API_KEY}",

            "Content-Type": "application/json"

        },

        json={

            "model": "mistralai/mistral-7b-instruct",

            "messages": [

                {

                    "role":"user",

                    "content": f"Summarize these study notes professionally with bullet points:\n\n{text}"

                }

            ]

        }

    )

    data = response.json()

    print(data)

    return data["choices"][0]["message"]["content"]


@app.route("/")
def home():

    return render_template("index.html")


@app.route("/summarize", methods=["POST"])
def summarize():

    try:

        notes = request.form.get("notes", "")

        uploaded_file = request.files.get("file")

        extracted_text = ""

        if uploaded_file:

            filename = uploaded_file.filename.lower()

            if filename.endswith(".txt"):

                extracted_text = uploaded_file.read().decode("utf-8")

            elif filename.endswith(".docx"):

                doc = Document(uploaded_file)

                for para in doc.paragraphs:
                    extracted_text += para.text + " "

            elif filename.endswith(".pdf"):

                extracted_text = extract_pdf(uploaded_file)

            elif filename.endswith(".pptx"):

                prs = Presentation(uploaded_file)

                for slide in prs.slides:

                    for shape in slide.shapes:

                        if hasattr(shape, "text"):
                            extracted_text += shape.text + " "

        final_text = notes + " " + extracted_text

        if final_text.strip() == "":

            return jsonify({
                "summary":"No text found."
            })

        summary = summarize_ai(final_text)

        return jsonify({
            "summary":summary
        })

    except Exception as e:

        return jsonify({
            "summary":str(e)
        })


if __name__ == "__main__":
    app.run(debug=True)